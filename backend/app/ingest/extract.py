"""
app/ingest/extract.py
~~~~~~~~~~~~~~~~~~~~~
Turn a file into an Extraction: text blocks (with page / time / heading), visual
items (photos, document images, scanned pages, video keyframes), the media file to
transcribe, and file metadata. No models run here; embeddings, faces, OCR and
transcription happen in the model stage (app/ingest/pipeline.py).

Memory stays flat for files of any size:
  * video keyframes are written to disk as they are decoded (never held in RAM)
  * scanned PDF pages are rendered lazily, one at a time, by the model stage
  * huge text files, CSV/TSV and XLSX sheets are streamed (lazy_blocks)

Two kinds of derived images are written:
  * display copies (small WebP) in the source's cache folder — served to the UI
  * work copies (JPEG, larger) in cache/_work — read by the models, then deleted
Original files are only ever read.
"""
from __future__ import annotations

import csv
import io
import math
import shutil
import subprocess
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, Iterator, List, Optional, Tuple

from PIL import Image, ImageOps

try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
except Exception:
    pass

from app.ingest.chunking import MAX_CHARS
from app.search.constants import LIBREOFFICE_EXTENSIONS, PLAIN_TEXT_EXTENSIONS

# Panoramas and large scans are fine; images far beyond any camera are refused instead
# of exhausting RAM (PIL raises above twice this many pixels).
Image.MAX_IMAGE_PIXELS = 250_000_000

DISPLAY_MAX_SIDE = 384          # thumbnails, keyframes, page previews served to the UI
DISPLAY_QUALITY = 75
KEYFRAME_WORK_SIDE = 1024       # faces need more resolution than the display copy has
DOC_IMAGE_WORK_SIDE = 1600      # document images keep enough detail for OCR
PAGE_RENDER_DPI = 200           # scanned pages, rendered for OCR
MIN_DOC_IMAGE_BYTES = 4096
MIN_DOC_IMAGE_SIDE = 64
MIN_PHOTO_SIDE = 128            # standalone images smaller than this (icons) are indexed by name only
MIN_PHOTO_SHORT_SIDE = 48       # ...and thinner ones (strips, separators, bars); panoramas stay well above
SCANNED_PAGE_MIN_CHARS = 40
KEYFRAME_MIN_INTERVAL = 2.0     # seconds between kept keyframes
KEYFRAME_MAX_GAP = 5.0          # fill gaps so there is at least one frame per 5 s
KEYFRAME_DEDUPE_BITS = 4        # dHash Hamming distance treated as "same shot"
KEYFRAME_FORCE_INTERVAL = 15.0  # ...but a static shot still gets a frame every 15 s
THUMB_CANDIDATES = 12           # the video thumbnail is the best of the first N keyframes
TEXT_PIECE_CHARS = 2_000_000
WORK_DIRNAME = "_work"


@dataclass
class TextBlock:
    text: str
    heading: Optional[str] = None
    page: Optional[int] = None
    kind: str = "body"                       # body | ocr | table (a table block is already one chunk)


@dataclass
class Visual:
    kind: str                                # photo | doc_image | page | keyframe
    path: Optional[Path] = None              # image the model stage loads (original photo or work copy)
    render: Optional[Callable[[], Image.Image]] = None   # lazily produced image (scanned pages)
    cache_rel: Optional[str] = None          # display copy inside the source cache folder
    write_cache: bool = False                # the model stage writes cache_rel from the loaded image
    work: bool = False                       # `path` is a temporary work copy, deleted after use
    page: Optional[int] = None
    ts: Optional[float] = None
    ocr: bool = True
    ocr_gate: bool = False                   # OCR only when the image looks like it contains text
    faces: bool = True
    page_text: bool = False                  # OCR text replaces a missing text layer
    width: Optional[int] = None
    height: Optional[int] = None


@dataclass
class Extraction:
    blocks: List[TextBlock] = field(default_factory=list)
    lazy_blocks: List[Callable[[], Iterator[TextBlock]]] = field(default_factory=list)
    sections: List[Dict[str, Any]] = field(default_factory=list)
    visuals: List[Visual] = field(default_factory=list)
    transcribe: Optional[Path] = None
    meta: Dict[str, Any] = field(default_factory=dict)
    log: List[str] = field(default_factory=list)
    thumbnail_rel: Optional[str] = None
    work_dir: Optional[Path] = None
    cleanup_hooks: List[Callable[[], None]] = field(default_factory=list)

    def iter_blocks(self) -> Iterator[TextBlock]:
        yield from self.blocks
        for produce in self.lazy_blocks:
            yield from produce()

    def cleanup(self) -> None:
        for hook in self.cleanup_hooks:
            try:
                hook()
            except Exception:
                pass
        self.cleanup_hooks.clear()
        if self.work_dir is not None:
            shutil.rmtree(self.work_dir, ignore_errors=True)


# ── image helpers ────────────────────────────────────────────────────────────
def open_image(src, max_side: Optional[int] = None) -> Image.Image:
    """Open as RGB with EXIF rotation applied; JPEGs are decoded at reduced scale when possible."""
    im = Image.open(io.BytesIO(src)) if isinstance(src, (bytes, bytearray)) else Image.open(str(src))
    with im:
        if getattr(im, "is_animated", False):
            im.seek(0)  # first frame of GIFs / animated WebP
        if max_side:
            try:
                im.draft("RGB", (max_side, max_side))
            except Exception:
                pass
        out = ImageOps.exif_transpose(im).convert("RGB")
    if max_side and max(out.size) > max_side:
        out.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
    return out


def image_size(path: Path) -> Tuple[int, int]:
    """Displayed size of an image file (EXIF rotation considered) without decoding it."""
    with Image.open(str(path)) as im:
        w, h = im.size
        try:
            if im.getexif().get(0x0112) in (5, 6, 7, 8):
                w, h = h, w
        except Exception:
            pass
    return w, h


def _shrunk(img: Image.Image, max_side: int) -> Image.Image:
    if max(img.size) <= max_side:
        return img
    im = img.copy()
    im.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
    return im


def save_display(img: Image.Image, dest: Path, max_side: int = DISPLAY_MAX_SIDE) -> Tuple[int, int]:
    dest.parent.mkdir(parents=True, exist_ok=True)
    im = _shrunk(img, max_side)
    im.save(str(dest), "WEBP", quality=DISPLAY_QUALITY, method=4)
    return im.width, im.height


def save_work(img: Image.Image, dest: Path, max_side: int, quality: int = 90) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    _shrunk(img, max_side).save(str(dest), "JPEG", quality=quality)
    return dest


def dhash_int(img: Image.Image) -> int:
    small = img.convert("L").resize((9, 8), Image.Resampling.BILINEAR)
    px = list(small.getdata())
    bits = 0
    for row in range(8):
        for col in range(8):
            bits = (bits << 1) | (1 if px[row * 9 + col] > px[row * 9 + col + 1] else 0)
    return bits


def _hamming(a: int, b: int) -> int:
    return bin(a ^ b).count("1")


def _exif_info(path: Path) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    try:
        with Image.open(str(path)) as im:
            exif = im.getexif()
            if not exif:
                return out
            sub = exif.get_ifd(0x8769) or {}
            raw = sub.get(36867) or exif.get(306)  # DateTimeOriginal, else DateTime
            if raw:
                try:
                    out["taken_at"] = datetime.strptime(str(raw).strip(), "%Y:%m:%d %H:%M:%S").isoformat()
                except ValueError:
                    pass
            gps = exif.get_ifd(0x8825) or {}
            if gps.get(2) and gps.get(4):
                def _deg(v):
                    d, m, s = (float(x) for x in v)
                    return d + m / 60 + s / 3600
                lat, lon = _deg(gps[2]), _deg(gps[4])
                if str(gps.get(1, "N")).upper().startswith("S"):
                    lat = -lat
                if str(gps.get(3, "E")).upper().startswith("W"):
                    lon = -lon
                if abs(lat) <= 90 and abs(lon) <= 180 and (lat or lon):
                    out["gps_lat"], out["gps_lon"] = round(lat, 6), round(lon, 6)
            if exif.get(0x0110):
                out["camera"] = str(exif.get(0x0110)).strip()
    except Exception:
        pass
    return out


# ── text helpers ─────────────────────────────────────────────────────────────
def _text_encoding(path: Path) -> str:
    with open(path, "rb") as f:
        head = f.read(65536)
    if head.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if head.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    try:
        head.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError as e:
        # A multi-byte character cut at the end of the sample is still UTF-8
        return "utf-8" if e.start >= len(head) - 3 else "cp1252"


def open_text(path: Path):
    return open(path, "r", encoding=_text_encoding(path), errors="replace", newline="")


def _iter_text_pieces(path: Path) -> Iterator[str]:
    """Read a text file of any size in pieces that end at line boundaries."""
    with open_text(path) as f:
        buf = ""
        while True:
            data = f.read(TEXT_PIECE_CHARS)
            if not data:
                break
            buf += data
            cut = buf.rfind("\n")
            if cut <= 0:
                if len(buf) < TEXT_PIECE_CHARS * 2:
                    continue
                cut = len(buf)  # one enormous line: cut anyway
            yield buf[:cut]
            buf = buf[cut + 1:]
        if buf.strip():
            yield buf


def _table_blocks(rows: Iterator[List[Any]], heading: str, ex: Extraction) -> Iterator[TextBlock]:
    """Stream rows as table chunks (first row = header, repeated in each chunk)."""
    header: Optional[str] = None
    current: List[str] = []
    size = 0
    chars = 0

    def block() -> TextBlock:
        nonlocal chars
        text = "\n".join(([header] if header else []) + current)
        chars += len(text)
        return TextBlock(text=text, heading=heading, kind="table")

    for row in rows:
        cells = ["" if v is None else str(v).strip() for v in row]
        line = " | ".join(c for c in cells if c)
        if not line:
            continue
        if header is None:
            header = line[:MAX_CHARS // 2]
            size = len(header)
            continue
        line = line[:MAX_CHARS]
        if current and size + len(line) + 1 > MAX_CHARS:
            yield block()
            current, size = [], len(header)
        current.append(line)
        size += len(line) + 1
    if current or header:
        # A sheet with only a header row is still searchable by its column names
        yield block()
    if chars:
        ex.sections.append({"title": heading, "page": None, "char_count": chars})


# ── documents ────────────────────────────────────────────────────────────────
class _PdfPages:
    """Renders PDF pages on demand (one open document, used by the model-stage thread)."""

    def __init__(self, path: Path):
        self.path = path
        self._doc = None

    def render(self, page_idx: int, dpi: int = PAGE_RENDER_DPI) -> Image.Image:
        import fitz
        if self._doc is None:
            self._doc = fitz.open(str(self.path))
        pix = self._doc[page_idx].get_pixmap(dpi=dpi)
        mode = "RGBA" if pix.alpha else "RGB"
        return Image.frombytes(mode, (pix.width, pix.height), pix.samples).convert("RGB")

    def close(self) -> None:
        if self._doc is not None:
            self._doc.close()
            self._doc = None


def _add_doc_image(ex: Extraction, cache_dir: Path, img: Optional[Image.Image], data_len: int,
                   idx: int, page: Optional[int]) -> None:
    if img is None or data_len < MIN_DOC_IMAGE_BYTES or min(img.width, img.height) < MIN_DOC_IMAGE_SIDE:
        return
    assert ex.work_dir is not None
    name = f"img_{idx:05d}"
    w, h = save_display(img, cache_dir / f"{name}.webp")
    work = save_work(img, ex.work_dir / f"{name}.jpg", DOC_IMAGE_WORK_SIDE, quality=92)
    ex.visuals.append(Visual(kind="doc_image", path=work, work=True, cache_rel=f"{name}.webp", page=page,
                             ocr_gate=True, width=w, height=h))


def _image_from_bytes(data: Optional[bytes]) -> Optional[Image.Image]:
    if not data:
        return None
    try:
        return open_image(data, max_side=DOC_IMAGE_WORK_SIDE)
    except Exception:
        return None


def _pdf_image(doc, xref: int) -> Tuple[Optional[Image.Image], int]:
    try:
        data = (doc.extract_image(xref) or {}).get("image")
    except Exception:
        data = None
    img = _image_from_bytes(data)
    if img is not None:
        return img, len(data or b"")
    # JBIG2 / JPX / CMYK and other encodings PIL can't open: let MuPDF decode them
    try:
        import fitz
        pix = fitz.Pixmap(doc, xref)
        if pix.n - pix.alpha >= 4:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        mode = "RGBA" if pix.alpha else "RGB"
        img = Image.frombytes(mode, (pix.width, pix.height), pix.samples).convert("RGB")
        return _shrunk(img, DOC_IMAGE_WORK_SIDE), len(data) if data else pix.width * pix.height
    except Exception:
        return None, 0


def _extract_pdf(path: Path, cache_dir: Path, ex: Extraction) -> None:
    import fitz
    from app.processing.processors import extract_page_heading_and_keywords

    pages = _PdfPages(path)
    ex.cleanup_hooks.append(pages.close)
    doc = fitz.open(str(path))
    try:
        ex.meta["page_count"] = len(doc)
        seen_xrefs = set()
        img_idx = 0
        scanned = 0
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            p_num = page_idx + 1
            text = page.get_text("text") or ""
            if len(text.strip()) < SCANNED_PAGE_MIN_CHARS:
                # No usable text layer: the model stage renders the page and its OCR text
                # becomes the page text (the page render covers any images on it)
                ex.visuals.append(Visual(
                    kind="page", render=lambda i=page_idx: pages.render(i),
                    cache_rel=f"page_{p_num:05d}.webp", write_cache=True, page=p_num, page_text=True,
                ))
                scanned += 1
                continue
            try:
                heading, _, _ = extract_page_heading_and_keywords(page, p_num, path.name)
            except Exception:
                heading = f"Page {p_num}"
            ex.blocks.append(TextBlock(text=text, heading=heading, page=p_num))
            ex.sections.append({"title": heading, "page": p_num, "char_count": len(text)})
            for info in page.get_images(full=True):
                xref = info[0]
                if xref in seen_xrefs:
                    continue  # repeated logos / headers
                seen_xrefs.add(xref)
                img, size = _pdf_image(doc, xref)
                if img is not None:
                    img_idx += 1
                    _add_doc_image(ex, cache_dir, img, size, img_idx, p_num)
        if scanned:
            ex.log.append(f"{scanned} page(s) without a text layer will be OCR'd")
        if len(doc):
            pix = doc[0].get_pixmap(dpi=60)
            mode = "RGBA" if pix.alpha else "RGB"
            img = Image.frombytes(mode, (pix.width, pix.height), pix.samples).convert("RGB")
            save_display(img, cache_dir / "thumb.webp")
            ex.thumbnail_rel = "thumb.webp"
    finally:
        doc.close()


def _docx_images(path: Path, cache_dir: Path, ex: Extraction) -> None:
    try:
        import docx
        d = docx.Document(str(path))
        idx = 0
        for rel in d.part.rels.values():
            if "image" in rel.reltype:
                blob = rel.target_part.blob
                idx += 1
                _add_doc_image(ex, cache_dir, _image_from_bytes(blob), len(blob), idx, None)
    except Exception as e:
        ex.log.append(f"DOCX image extraction skipped: {e}")


def _pptx_images(path: Path, cache_dir: Path, ex: Extraction) -> None:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(str(path))
        idx = 0
        for s_idx, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    blob = shape.image.blob
                    idx += 1
                    _add_doc_image(ex, cache_dir, _image_from_bytes(blob), len(blob), idx, s_idx)
    except Exception as e:
        ex.log.append(f"PPTX image extraction skipped: {e}")


def _xlsx_rows(path: Path, sheet_name: str) -> Iterator[List[Any]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        ws = wb[sheet_name]
        try:
            ws.reset_dimensions()  # some writers store a wrong sheet size; read every row present
        except Exception:
            pass
        for row in ws.iter_rows(values_only=True):
            yield list(row)
    finally:
        wb.close()


def _delimited_rows(path: Path, delimiter: str) -> Iterator[List[Any]]:
    csv.field_size_limit(1 << 30)
    with open_text(path) as f:
        for row in csv.reader(f, delimiter=delimiter):
            yield row


def extract_document(path: Path, ext: str, cache_dir: Path, ex: Extraction) -> Extraction:
    from app.processing.processors import (
        _convert_with_libreoffice, _extract_docx_sections, _extract_pptx_sections, _read_text_file,
    )
    ext = ext.lower()

    if ext == "pdf":
        _extract_pdf(path, cache_dir, ex)

    elif ext == "docx" or ext in LIBREOFFICE_EXTENSIONS:
        docx_path, tmp_dir = str(path), None
        if ext != "docx":
            docx_path, tmp_dir = _convert_with_libreoffice(str(path), "docx")
            if not docx_path:
                raise RuntimeError(f"Cannot read .{ext} files without LibreOffice; install it or save as .docx")
        try:
            sections, _, _ = _extract_docx_sections(docx_path)
            for s in sections:
                ex.blocks.append(TextBlock(text=s["section_text"], heading=s["section_title"], page=s.get("page_num")))
                ex.sections.append({"title": s["section_title"], "page": s.get("page_num"),
                                    "char_count": len(s["section_text"])})
            _docx_images(Path(docx_path), cache_dir, ex)
        finally:
            if tmp_dir:
                shutil.rmtree(tmp_dir, ignore_errors=True)

    elif ext == "pptx":
        sections, _ = _extract_pptx_sections(str(path))
        ex.meta["page_count"] = len(sections)
        for s in sections:
            ex.blocks.append(TextBlock(text=s["section_text"], heading=s["section_title"], page=s["page_num"]))
            ex.sections.append({"title": s["section_title"], "page": s["page_num"], "char_count": len(s["section_text"])})
        _pptx_images(path, cache_dir, ex)

    elif ext == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True)
        try:
            names = list(wb.sheetnames)
        finally:
            wb.close()
        for name in names:
            ex.lazy_blocks.append(lambda n=name: _table_blocks(_xlsx_rows(path, n), f"Sheet: {n}", ex))

    elif ext == "xls":
        import xlrd
        book = xlrd.open_workbook(str(path))
        for sheet in book.sheets():
            rows = [sheet.row_values(r) for r in range(sheet.nrows)]
            ex.blocks.extend(_table_blocks(iter(rows), f"Sheet: {sheet.name}", ex))

    elif ext in ("csv", "tsv"):
        delimiter = "\t" if ext == "tsv" else ","
        ex.lazy_blocks.append(lambda: _table_blocks(_delimited_rows(path, delimiter), path.stem, ex))

    elif ext in ("html", "htm"):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(_read_text_file(str(path)), "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        title = soup.find("title")
        text = "\n".join(line for line in (l.strip() for l in soup.get_text("\n").splitlines()) if line)
        heading = title.get_text(strip=True) if title else path.stem
        ex.blocks.append(TextBlock(text=text, heading=heading))
        ex.sections.append({"title": heading, "page": None, "char_count": len(text)})

    elif ext in PLAIN_TEXT_EXTENSIONS:
        def _plain() -> Iterator[TextBlock]:
            total = 0
            for piece in _iter_text_pieces(path):
                total += len(piece)
                yield TextBlock(text=piece, heading=path.stem)
            ex.sections.append({"title": path.stem, "page": None, "char_count": total})
        ex.lazy_blocks.append(_plain)

    else:
        raise RuntimeError(f"Unsupported document type '.{ext}'")
    return ex


# ── images ───────────────────────────────────────────────────────────────────
def is_undersized_photo(w: int, h: int) -> bool:
    return max(w, h) < MIN_PHOTO_SIDE or min(w, h) < MIN_PHOTO_SHORT_SIDE


def extract_image(path: Path, cache_dir: Path, ex: Extraction) -> Extraction:
    w, h = image_size(path)
    ex.meta.update({"width": w, "height": h})
    ex.meta.update(_exif_info(path))
    save_display(open_image(path, max_side=DISPLAY_MAX_SIDE * 2), cache_dir / "thumb.webp")
    ex.thumbnail_rel = "thumb.webp"
    if is_undersized_photo(w, h):
        # An icon, button or strip graphic, not a photo: still findable by name, but no
        # image vector, OCR, faces or description (they would only add noise to results)
        ex.log.append(f"{w}x{h} px: too small to be a photo, indexed by name only")
        ex.meta["small_image"] = True
        return ex
    # The model stage reads the original (decoded at reduced scale)
    ex.visuals.append(Visual(kind="photo", path=path, cache_rel="thumb.webp", ocr_gate=True, width=w, height=h))
    return ex


# ── audio / video ────────────────────────────────────────────────────────────
def _probe(path: Path) -> Dict[str, Any]:
    info: Dict[str, Any] = {}
    try:
        import av
        with av.open(str(path)) as c:
            if c.duration:
                info["duration"] = float(c.duration) / 1_000_000
            video = next((s for s in c.streams if s.type == "video"), None)
            info["has_audio"] = any(s.type == "audio" for s in c.streams)
            if video is not None:
                info["width"] = video.codec_context.width
                info["height"] = video.codec_context.height
                if "duration" not in info and video.duration and video.time_base:
                    info["duration"] = float(video.duration * video.time_base)
                if not video.frames and video.average_rate and info.get("duration"):
                    info["frames"] = int(info["duration"] * float(video.average_rate))
                else:
                    info["frames"] = int(video.frames or 0)
            elif "duration" not in info:
                for s in c.streams:
                    if s.duration and s.time_base:
                        info["duration"] = float(s.duration * s.time_base)
                        break
    except Exception:
        pass
    if not info.get("duration"):
        try:
            out = subprocess.run(
                ["ffprobe", "-v", "error", "-show_entries", "format=duration",
                 "-of", "default=noprint_wrappers=1:nokey=1", str(path)],
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120,
            )
            info["duration"] = float(out.stdout.strip())
        except Exception:
            pass
    return info


def _frame_image(frame) -> Image.Image:
    """Decode a video frame to RGB at most KEYFRAME_WORK_SIDE, honouring the rotation flag."""
    w, h = frame.width, frame.height
    scale = min(1.0, KEYFRAME_WORK_SIDE / float(max(w, h)))
    tw, th = max(2, int(w * scale) // 2 * 2), max(2, int(h * scale) // 2 * 2)
    img = frame.to_image(width=tw, height=th)
    rotation = int(round(getattr(frame, "rotation", 0) or 0)) % 360
    if rotation == 90:
        img = img.transpose(Image.Transpose.ROTATE_90)
    elif rotation == 180:
        img = img.transpose(Image.Transpose.ROTATE_180)
    elif rotation == 270:
        img = img.transpose(Image.Transpose.ROTATE_270)
    return img.convert("RGB")


def _thumbnail_score(img: Image.Image) -> float:
    import numpy as np
    small = img.convert("L")
    small.thumbnail((160, 160))
    arr = np.asarray(small, dtype=np.float32)
    brightness = float(arr.mean())
    if brightness < 15 or brightness > 245:
        return -1.0
    return float(arr.var()) * (1.0 - abs(brightness - 128) / 256.0)


class _Keyframes:
    """Collects keyframes straight to disk; only their metadata stays in memory."""

    def __init__(self, cache_dir: Path, work_dir: Path):
        self.cache_dir = cache_dir
        self.work_dir = work_dir
        self.items: List[Dict[str, Any]] = []
        self.sampled = 0
        self._names: set = set()
        self._last: Optional[Tuple[float, int]] = None   # (ts, hash) of the last frame kept in this pass

    @staticmethod
    def _same_shot(ts: float, h: int, last: Optional[Tuple[float, int]]) -> bool:
        # A long static shot still gets a frame every KEYFRAME_FORCE_INTERVAL seconds
        return (last is not None and _hamming(h, last[1]) <= KEYFRAME_DEDUPE_BITS
                and ts - last[0] < KEYFRAME_FORCE_INTERVAL)

    def offer(self, ts: float, img: Image.Image, dedupe: bool = True) -> None:
        self.sampled += 1
        h = dhash_int(img)
        if dedupe and self._same_shot(ts, h, self._last):
            return
        self._last = (ts, h)
        name = f"kf_{int(round(ts * 1000)):010d}"
        if name in self._names:
            return
        self._names.add(name)
        w, hh = save_display(img, self.cache_dir / f"{name}.webp")
        save_work(img, self.work_dir / f"{name}.jpg", KEYFRAME_WORK_SIDE)
        self.items.append({"ts": ts, "hash": h, "name": name, "w": w, "h": hh, "score": _thumbnail_score(img)})

    def reset(self) -> None:
        for it in self.items:
            self._delete(it)
        self.items, self._names, self._last, self.sampled = [], set(), None, 0

    def _delete(self, it: Dict[str, Any]) -> None:
        (self.cache_dir / f"{it['name']}.webp").unlink(missing_ok=True)
        (self.work_dir / f"{it['name']}.jpg").unlink(missing_ok=True)

    def finish(self, ex: Extraction) -> None:
        self.items.sort(key=lambda it: it["ts"])
        kept: List[Dict[str, Any]] = []
        last: Optional[Tuple[float, int]] = None
        for it in self.items:
            if self._same_shot(it["ts"], it["hash"], last):
                self._delete(it)  # gap-fill frame of a shot we already have
                continue
            kept.append(it)
            last = (it["ts"], it["hash"])
        for it in kept:
            ex.visuals.append(Visual(
                kind="keyframe", path=self.work_dir / f"{it['name']}.jpg", work=True,
                cache_rel=f"{it['name']}.webp", ts=round(it["ts"], 2), ocr=False, width=it["w"], height=it["h"],
            ))
        candidates = kept[:THUMB_CANDIDATES]
        if candidates:
            best = max(candidates, key=lambda it: it["score"])
            shutil.copyfile(self.cache_dir / f"{best['name']}.webp", self.cache_dir / "thumb.webp")
            ex.thumbnail_rel = "thumb.webp"
        ex.log.append(f"{len(kept)} keyframe(s) kept from {self.sampled} sampled")


def _decode_keyframes(path: Path, frames: _Keyframes) -> Tuple[List[float], int]:
    """Pass 1: key frames only (cheap). Returns (sampled timestamps, number of key frames seen)."""
    import av
    sampled: List[float] = []
    seen = 0
    with av.open(str(path)) as container:
        stream = next((s for s in container.streams if s.type == "video"), None)
        if stream is None:
            return sampled, 0
        stream.codec_context.skip_frame = "NONKEY"
        stream.thread_type = "AUTO"
        last_ts = -1e9
        for frame in container.decode(stream):
            if frame.time is None:
                continue
            seen += 1
            ts = float(frame.time)
            if ts - last_ts < KEYFRAME_MIN_INTERVAL:
                continue
            last_ts = ts
            sampled.append(ts)
            frames.offer(ts, _frame_image(frame))
    return sampled, seen


def _frames_by_seeking(path: Path, wanted: List[float], frames: _Keyframes) -> None:
    import av
    with av.open(str(path)) as container:
        stream = next((s for s in container.streams if s.type == "video"), None)
        if stream is None or not stream.time_base:
            return
        stream.thread_type = "AUTO"
        for t in wanted:
            try:
                container.seek(int(t / stream.time_base), stream=stream, backward=True)
                for frame in container.decode(stream):
                    if frame.time is not None and float(frame.time) >= t - 0.05:
                        frames.offer(float(frame.time), _frame_image(frame), dedupe=False)
                        break
            except Exception:
                continue


def _frames_sequential(path: Path, wanted: List[float], frames: _Keyframes) -> None:
    """One decoding pass, picking the first frame at or after each wanted timestamp."""
    import av
    targets = sorted(wanted)
    i = 0
    with av.open(str(path)) as container:
        stream = next((s for s in container.streams if s.type == "video"), None)
        if stream is None:
            return
        stream.thread_type = "AUTO"
        for frame in container.decode(stream):
            if i >= len(targets):
                break
            if frame.time is None or float(frame.time) < targets[i] - 0.05:
                continue
            frames.offer(float(frame.time), _frame_image(frame), dedupe=False)
            while i < len(targets) and targets[i] - 0.05 <= float(frame.time):
                i += 1


def _frames_ffmpeg(path: Path, tmp_dir: Path, duration: Optional[float], frames: _Keyframes) -> None:
    """Fallback when PyAV cannot decode: one frame every KEYFRAME_MAX_GAP seconds via ffmpeg."""
    tmp_dir.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run(
            ["ffmpeg", "-y", "-loglevel", "error", "-i", str(path), "-vf",
             f"fps=1/{KEYFRAME_MAX_GAP},scale='min({KEYFRAME_WORK_SIDE},iw)':-2", "-q:v", "3",
             str(tmp_dir / "ff_%06d.jpg")],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=max(600, int((duration or 0) * 2)),
        )
        for i, f in enumerate(sorted(tmp_dir.glob("ff_*.jpg"))):
            with Image.open(f) as im:
                frames.offer(i * KEYFRAME_MAX_GAP, im.convert("RGB"))
            f.unlink(missing_ok=True)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def extract_video(path: Path, cache_dir: Path, ex: Extraction) -> Extraction:
    assert ex.work_dir is not None
    info = _probe(path)
    duration = info.get("duration")
    ex.meta["duration"] = duration
    if info.get("width"):
        ex.meta["width"], ex.meta["height"] = info["width"], info["height"]
    if info.get("has_audio", True):
        ex.transcribe = path

    frames = _Keyframes(cache_dir, ex.work_dir)
    try:
        sampled, keyframes_seen = _decode_keyframes(path, frames)
        if duration:
            bounds = [0.0] + sorted(sampled) + [duration]
            wanted: List[float] = []
            for a, b in zip(bounds, bounds[1:]):
                gap = b - a
                if gap > KEYFRAME_MAX_GAP:
                    n = int(math.ceil(gap / KEYFRAME_MAX_GAP)) - 1
                    wanted.extend(a + (k + 1) * gap / (n + 1) for k in range(n))
            if wanted:
                # Each seek decodes on average half a GOP; with sparse key frames one
                # sequential pass is cheaper than many seeks
                if len(wanted) < 2 * max(1, keyframes_seen):
                    _frames_by_seeking(path, wanted, frames)
                else:
                    _frames_sequential(path, wanted, frames)
    except Exception as e:
        ex.log.append(f"PyAV keyframes failed ({e}); using ffmpeg sampling")
        frames.reset()
        _frames_ffmpeg(path, ex.work_dir / "ffmpeg", duration, frames)
    frames.finish(ex)

    if ex.visuals and not ex.meta.get("width"):
        ex.meta["width"], ex.meta["height"] = ex.visuals[0].width, ex.visuals[0].height
    return ex


def extract_audio(path: Path, cache_dir: Path, ex: Extraction) -> Extraction:
    ex.meta["duration"] = _probe(path).get("duration")
    ex.transcribe = path
    return ex


EXTRACTORS: Dict[str, Callable[[Path, str, Path, Extraction], Extraction]] = {
    "image": lambda path, ext, cache, ex: extract_image(path, cache, ex),
    "video": lambda path, ext, cache, ex: extract_video(path, cache, ex),
    "audio": lambda path, ext, cache, ex: extract_audio(path, cache, ex),
    "document": lambda path, ext, cache, ex: extract_document(path, ext, cache, ex),
}


def extract(kind: str, path: Path, ext: str, cache_dir: Path) -> Extraction:
    """Extract a file. On failure the partial work folder is removed before re-raising."""
    if kind not in EXTRACTORS:
        raise RuntimeError(f"Unsupported file kind '{kind}'")
    cache_dir.mkdir(parents=True, exist_ok=True)
    ex = Extraction(work_dir=cache_dir / WORK_DIRNAME)
    ex.work_dir.mkdir(parents=True, exist_ok=True)
    try:
        return EXTRACTORS[kind](path, ext, cache_dir, ex)
    except BaseException:
        ex.cleanup()
        raise
