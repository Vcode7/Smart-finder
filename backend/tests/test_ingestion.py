"""
Ingestion regression tests (isolated temp database, fake models).
"""
import uuid
from pathlib import Path

import pytest

from app.database.session import db_all, db_get, db_run


def _insert_source(path: Path, file_type: str, status: str = "uploaded", attempts: int = 0) -> str:
    db_run(
        "INSERT OR IGNORE INTO users (id, email, username, password_hash, role) VALUES ('u1', 'u@x', 'u', 'x', 'user')"
    )
    sid = str(uuid.uuid4())
    db_run(
        """INSERT INTO knowledge_sources
           (id, filename, original_name, file_type, file_path, file_size, uploaded_by, processing_status, attempts)
           VALUES (?, ?, ?, ?, ?, ?, 'u1', ?, ?)""",
        (sid, path.name, path.name, file_type, str(path), path.stat().st_size if path.exists() else 0, status, attempts)
    )
    return sid


async def test_reprocessing_is_idempotent(tmp_db, tmp_faiss, mock_models_available, tmp_path):
    """Regression: an interrupted or repeated run duplicated sections, chunks and vectors."""
    from app.processing.queue import _process_source
    doc = tmp_path / "plants.txt"
    doc.write_text("Photosynthesis converts light energy into chemical energy. " * 60, encoding="utf-8")
    sid = _insert_source(doc, "txt")

    await _process_source(sid)
    chunks_first = len(db_all("SELECT id FROM document_chunks WHERE source_id = ?", (sid,)))
    vectors_first = sum(1 for m in tmp_faiss.text_meta if m["source_id"] == sid)
    assert chunks_first > 0 and vectors_first == chunks_first

    await _process_source(sid)
    assert len(db_all("SELECT id FROM document_chunks WHERE source_id = ?", (sid,))) == chunks_first
    assert len(db_all("SELECT id FROM document_sections WHERE source_id = ?", (sid,))) == 1
    assert sum(1 for m in tmp_faiss.text_meta if m["source_id"] == sid) == vectors_first
    assert db_get("SELECT processing_status FROM knowledge_sources WHERE id = ?", (sid,))["processing_status"] == "completed"


async def test_unsupported_type_is_marked_failed(tmp_db, tmp_faiss, tmp_path):
    from app.processing.queue import _process_source
    f = tmp_path / "program.exe"
    f.write_bytes(b"MZ\x00\x00")
    sid = _insert_source(f, "exe")
    await _process_source(sid)
    row = db_get("SELECT processing_status, error_message FROM knowledge_sources WHERE id = ?", (sid,))
    assert row["processing_status"] == "failed"
    assert "Unsupported" in row["error_message"]


def test_stale_sources_are_requeued_then_given_up(tmp_db, tmp_path):
    from app.processing.queue import recover_stale_sources, MAX_ATTEMPTS
    f = tmp_path / "a.txt"
    f.write_text("x")
    retry = _insert_source(f, "txt", status="processing", attempts=1)
    give_up = _insert_source(f, "txt", status="indexing", attempts=MAX_ATTEMPTS)
    recover_stale_sources()
    assert db_get("SELECT processing_status FROM knowledge_sources WHERE id = ?", (retry,))["processing_status"] == "uploaded"
    assert db_get("SELECT processing_status FROM knowledge_sources WHERE id = ?", (give_up,))["processing_status"] == "failed"


def test_image_without_clip_stores_no_nan_embedding(tmp_db, tmp_faiss, mock_models_available, monkeypatch, tmp_path):
    """Regression: float_array_to_blob(None) stored a NaN blob when CLIP was unavailable."""
    from PIL import Image
    import app.processing.processors as processors
    monkeypatch.setattr(processors, "get_image_embedding", lambda *_a, **_k: None)
    img = tmp_path / "photo.png"
    Image.new("RGB", (64, 48), (200, 30, 30)).save(img)
    sid = _insert_source(img, "png")
    processors.process_image(sid)
    row = db_get("SELECT embedding FROM images WHERE source_id = ?", (sid,))
    assert row is not None and row["embedding"] is None
    assert tmp_faiss.image_index.ntotal == 0


def test_docx_tables_are_extracted_in_document_order(tmp_path):
    import docx
    from app.processing.processors import _extract_docx_sections
    d = docx.Document()
    d.add_heading("Budget", level=1)
    d.add_paragraph("The table below lists costs.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Cost"
    table.cell(1, 0).text, table.cell(1, 1).text = "Laptop", "900"
    d.add_paragraph("Totals follow the table.")
    path = tmp_path / "budget.docx"
    d.save(path)

    sections, full_text, _ = _extract_docx_sections(str(path))
    assert "Laptop | 900" in full_text
    assert full_text.index("lists costs") < full_text.index("Laptop | 900") < full_text.index("Totals follow")


def test_pptx_slides_and_notes(tmp_path):
    from pptx import Presentation
    from app.processing.processors import _extract_pptx_sections
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Results"
    slide.placeholders[1].text = "Revenue grew 12 percent"
    slide.notes_slide.notes_text_frame.text = "Mention the new market"
    path = tmp_path / "deck.pptx"
    prs.save(path)

    sections, full_text = _extract_pptx_sections(str(path))
    assert sections[0]["section_title"].startswith("Slide 1")
    assert "Revenue grew 12 percent" in full_text
    assert "Speaker notes: Mention the new market" in full_text


def test_utf16_text_files_are_decoded(tmp_path):
    from app.processing.processors import _read_text_file
    path = tmp_path / "notes.txt"
    path.write_text("नमस्ते and hello", encoding="utf-16")
    assert _read_text_file(str(path)) == "नमस्ते and hello"


def test_sanitize_filename():
    from app.core.files import sanitize_filename, sanitize_relative_path
    assert sanitize_filename("../../etc/passwd") == "passwd"
    assert sanitize_filename("C:\\Windows\\evil.bat") == "evil.bat"
    assert sanitize_filename('a<b>c:"d|e?.txt') == "a_b_c__d_e_.txt"
    assert sanitize_filename("CON.txt") == "_CON.txt"
    assert sanitize_filename("...") == "file"
    assert str(sanitize_relative_path("../a/../b/c.txt")).replace("\\", "/") == "a/b/c.txt"
